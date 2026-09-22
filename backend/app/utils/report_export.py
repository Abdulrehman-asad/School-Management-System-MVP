"""
Generic report export helpers: turn a list of dicts (+ column headers) into
either an .xlsx workbook or a .pdf document, both returned as in-memory bytes
ready to stream back as a file download.

Used by app/controllers/report_controller.py for Attendance, Result, Student,
Teacher, and Fee reports.
"""

from io import BytesIO
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def build_excel_report(title: str, headers: List[str], rows: List[Dict[str, Any]]) -> BytesIO:
    """rows: list of dicts whose keys match `headers` (case-sensitive)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else "Report"  # Excel sheet names cap at 31 chars

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(size=14, bold=True)
    title_cell.alignment = Alignment(horizontal="center")

    # Header row
    header_row_idx = 3
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_offset, row_data in enumerate(rows, start=1):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=header_row_idx + row_offset, column=col_idx, value=row_data.get(header, ""))

    # Auto-fit column widths (approximate, based on content length)
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_data in rows:
            value = row_data.get(header, "")
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_pdf_report(title: str, headers: List[str], rows: List[Dict[str, Any]],
                      subtitle: str = "") -> BytesIO:
    buffer = BytesIO()
    page_size = landscape(A4) if len(headers) > 5 else A4

    doc = SimpleDocTemplate(
        buffer, pagesize=page_size,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Shaheen Model Girls High School", styles["Title"]))
    elements.append(Paragraph(title, styles["Heading2"]))
    if subtitle:
        elements.append(Paragraph(subtitle, styles["Normal"]))
    elements.append(Spacer(1, 0.5 * cm))

    table_data = [headers] + [[str(row.get(h, "")) for h in headers] for row in rows]

    if len(rows) == 0:
        elements.append(Paragraph("No records found for the selected filters.", styles["Normal"]))
    else:
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ]))
        elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer
